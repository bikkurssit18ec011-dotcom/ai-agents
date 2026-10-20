import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

OUT = sys.argv[1]

# Dezmos brand tokens
LIME, FOREST, GRAY, TINT, WHITE = "C1FF72", "1F2A2E", "F4F8FA", "F3FFE3", "FFFFFF"
MUTED = "5A6366"
FONT = "Manrope"
f_title = Font(name=FONT, size=18, bold=True, color=FOREST)
f_sub = Font(name=FONT, size=10, color=MUTED)
f_head = Font(name=FONT, size=10, bold=True, color=WHITE)
f_body = Font(name=FONT, size=10, color=FOREST)
f_bold = Font(name=FONT, size=10, bold=True, color=FOREST)
f_kpi = Font(name=FONT, size=22, bold=True, color=FOREST)
f_link = Font(name=FONT, size=10, color=FOREST, underline="single")
fill_head = PatternFill("solid", fgColor=FOREST)
fill_lime = PatternFill("solid", fgColor=LIME)
fill_gray = PatternFill("solid", fgColor=GRAY)
fill_tint = PatternFill("solid", fgColor=TINT)
thin = Side(style="thin", color="D9DDDE")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

NICHES = ["Solar Installation", "Spice Brand", "Food Manufacturer", "Interior Design", "Furniture Shop"]
STATUSES = ["Not Contacted", "Called - No Answer", "Called - Call Back Later", "Interested",
            "Meeting / Proposal Sent", "Converted", "Not Interested", "Invalid Number"]
OUTCOMES = ["No Answer", "Call Back Later", "Interested", "Meeting Booked", "Proposal Sent",
            "Converted", "Not Interested", "Wrong / Invalid Number"]

JD = "Justdial virtual no. (routes to owner)"
IM = "IndiaMART virtual no. (routes to owner)"
MOB = "Direct mobile"
LAND = "Office landline"
NOWEB = "No website found"
WEB = "Has website - audit conversions"

# Source URLs
S = {
    "jd_solar": "https://www.justdial.com/Thiruvananthapuram/Residential-Solar-Panel-Installation-Services/nct-12154454",
    "enf": "https://www.enfsolar.com/",
    "kseb": "https://ekiran.kseb.in/documents/2%20and%20upto%203.pdf",
    "im_spice": "https://dir.indiamart.com/thiruvananthapuram/spices.html",
    "im_snack": "https://dir.indiamart.com/thiruvananthapuram/banana-chips.html",
    "im_bakery": "https://dir.indiamart.com/thiruvananthapuram/bakery-products.html",
    "dnb_food": "https://www.dnb.com/business-directory/company-information.food_manufacturing.in.kerala.thiruvananthapuram.html",
    "jd_int": "https://www.justdial.com/Thiruvananthapuram/Interior-Designers-in-Kazhakuttam/nct-10272436",
    "gkw": "https://interiordesign.gkwretail.com/interior-designers-in-trivandrum/",
    "jd_furn": "https://www.justdial.com/Thiruvananthapuram/Wooden-Furniture-Dealers/nct-10542688",
    "yappe": "https://yappe.in/kerala/thiruvananthapuram/",
    "qk_furn": "https://www.quickerala.com/trivandrum/furniture-and-fixtures/furniture-shop/sbct-8435-dt-12",
}

# niche, name, owner, area, address, phone, phone_type, alt, email, website, rating, source, url, hook
L = [
 # ---------------- SOLAR ----------------
 ("Solar Installation","Aadithya Innovations Solar","","Pappanamcode","TC 56/60, 10, SRK Building, Pappanamcode, TVM 695018","08490040914",JD,"","","aadithyainnovations.com","4.9 (129) Justdial","Justdial + ENF Solar",S["jd_solar"],"Strong 4.9 rating with 129 reviews - showcase it on the site to convert visitors"),
 ("Solar Installation","Anand Solar (TATA Power Solar Distributor)","","Peyad","Peyad, Thiruvananthapuram","08128463201",JD,"","","","4.9 Justdial","Justdial",S["jd_solar"],"TATA Power brand authority but no own website - competitors outrank them on Google"),
 ("Solar Installation","Micro Smart Technologies","","Peyad","Peyad, Thiruvananthapuram","08460307105",JD,"","","","","Justdial",S["jd_solar"],""),
 ("Solar Installation","Hello Solar Energy Solutions","","Varkala","Varkala, Thiruvananthapuram district","08460530606",JD,"","","","","Justdial",S["jd_solar"],"Varkala homestays/resorts are a high-ticket solar segment - landing page for them"),
 ("Solar Installation","V Power Solutions","","Attukal","Attukal, Thiruvananthapuram","09972269871",JD,"","","","","Justdial",S["jd_solar"],""),
 ("Solar Installation","Rudraveena Solar & Home Solutions","","Naruvamoodu","Naruvamoodu, Thiruvananthapuram","+91 85900 02123",MOB,"09606167175 (Justdial)","contact@rudraveenasolutions.com","rudraveenasolutions.com","","Website + Justdial","https://www.rudraveenasolutions.com/","Claims KSEB empanelment - make subsidy eligibility the hero of the site"),
 ("Solar Installation","Univarsa Techies Pvt Ltd","","Enjakkal","Enjakkal, Thiruvananthapuram","09008359088",JD,"","","","","Justdial",S["jd_solar"],""),
 ("Solar Installation","Light Sky International Pvt Ltd","","Jagathy","Sree Rudram, Jagathy, TVM 695014","+91 73562 13562",MOB,"08460379701 (Justdial)","info@lightskyinternational.com","lightskyinternational.com","","Website + Justdial","https://lightskyinternational.com/",""),
 ("Solar Installation","Wattsun Energy","","Chackai","Airport Road, Chackai, Thiruvananthapuram","+91 93881 06363",MOB,"+91 89433 45706","sourawattsun@gmail.com","","","KSEB Soura empanelled list",S["kseb"],"KSEB Soura empanelled but only a Gmail ID - no website to capture subsidy searches"),
 ("Solar Installation","Renergy Systems","","Attingal","Attingal, Thiruvananthapuram district","+91 70343 22221",MOB,"+91 96453 22221","soura.renergysales@gmail.com","","","KSEB Soura empanelled list",S["kseb"],"KSEB Soura empanelled but only a Gmail ID - no website"),
 ("Solar Installation","BSS Solar Kerala","","Kowdiar","Brahmins Colony, Kowdiar PO, TVM 695003","0471-2439322",LAND,"9048001332","bsssolarkerala@gmail.com","bsssolar.com","","Company website","https://bsssolar.com/",""),
 ("Solar Installation","OnGrid Solar Power Solutions Pvt Ltd","","Vattiyoorkavu","Vattiyoorkavu, Thiruvananthapuram","+91 75949 49406",MOB,"","ongridsolarpowersolution@gmail.com","ongridsolarpowersolutions.com","","Company website","https://www.ongridsolarpowersolutions.com/","Advertises up to Rs 78,000 subsidy - build a subsidy calculator funnel around it"),
 ("Solar Installation","Green Factory","","Fort","Fort PO, Thiruvananthapuram 695023","+91 89433 45715",MOB,"","info@greenfactoryindia.com","greenfactoryindia.com","","Company website","https://greenfactoryindia.com/contact-us",""),
 ("Solar Installation","SEnSOL Technologies","","Pazhavangadi","Pazhavangadi, Fort PO, TVM 695023","+91 94472 22800",MOB,"+91 94472 22900","sensolindia@gmail.com","sensol.in","","Company website","https://sensol.in/contact-us/","20+ years in solar PV - case studies with kW + savings would prove it"),
 ("Solar Installation","VAIGAA Power Solutions","","Vembayam","Near Indian Oil petrol pump, Vembayam, TVM","+91 85906 15993",MOB,"+91 70128 01008 / +91 73564 00461","vaigaapowersolutions@gmail.com","vaigaapowersolutions.in","","Company website","https://vaigaapowersolutions.in/",""),
 ("Solar Installation","Comorin Solar","","Ulloor / Technopark","Kochulloor Junction, Ulloor / Technopark Phase-3","+91 96331 59523",MOB,"","comorinsolar@gmail.com","comorinsolar.com","","Website + Justdial","https://www.comorinsolar.com/",""),
 ("Solar Installation","Chirayil Solar Power Tech","","Nalanchira","Nalanchira, Thiruvananthapuram 695015","+91 97788 05529",MOB,"","chirayilpower@gmail.com","","","Web search listing","https://vguardsolartvm.com/","Website not confirmed - check before call"),
 ("Solar Installation","Wega Solar (Wegasolar Renewables)","","Sreekariyam","Sreekariyam, Thiruvananthapuram 695017","+91 62388 10541",MOB,"","info@wegasolar.in","wegasolar.in","","Company website","https://www.wegasolar.in/",""),
 ("Solar Installation","Mitra Power Solutions","","Kumarapuram","TC 95/324-7, KIMS Road, Kumarapuram, TVM 695011","",""," ","","mitrapower.in","","ENF Solar directory","https://www.enfsolar.com/mitra-power-solutions","Phone not public - get it from the website contact page"),
 ("Solar Installation","Hive Solar Power Solution","","Vazhuthacaud","Opp Cotton Hill LP School, Vazhuthacaud, TVM 695014","","","","","hivesolar.com","","ENF Solar + Justdial","https://www.enfsolar.com/hive-solar","Installing since 2013 - years of projects to turn into case studies"),
 ("Solar Installation","Ray Technologies","","Mannanthala","Nalanchira, Mannanthala, TVM 695015","","","","","raytechnologiestvm.com","","ENF Solar directory","https://www.enfsolar.com/ray-technologies",""),
 ("Solar Installation","Premier Projects","","Peroorkada","Muthassi Towers, Mannamoola, Peroorkada, TVM 695005","","","","","premierprojects.in","","ENF Solar directory","https://www.enfsolar.com/premier-projects",""),
 ("Solar Installation","Smart Solar Global / Smart Solar Homes","","Nalanchira","Thiruvathira, Kottamughal, Nalanchira PO, TVM 695015","","","","","smartsolarglobal.com","","ENF Solar + Justdial","https://www.enfsolar.com/smart-solar-5","1 MWp+ projects - a commercial/industrial landing page fits"),
 ("Solar Installation","Green Sun","","Thiruvallom","Crescent Hill, Thiruvallom, TVM 695027","","","","","greensun.in","","ENF Solar directory","https://www.enfsolar.com/green-sun-2",""),
 ("Solar Installation","Enegix Enterprises","","Sasthamangalam","TC 9-2769, CSM Nagar 286, Edapazhanji, Sasthamangalam, TVM 695010","","","","","enegix.co.in","","ENF Solar directory","https://www.enfsolar.com/enegix-enterprises",""),
 ("Solar Installation","El Sol Power Solutions Pvt Ltd","Tinsu Mathew (Managing Director)","Trivandrum","Thiruvananthapuram","","","","","elsol.co.in","","Company website","https://elsol.co.in/locations/solar-company-in-trivandrum/","Named MD - ask for Tinsu directly"),
 # ---------------- SPICE ----------------
 ("Spice Brand","Azad Spice Gardens Pvt Ltd","Ozman Azad (Proprietor, per TradeIndia)","Kowdiar","TC-4/1073-4, Kuravankonam Road, Kowdiar, TVM 695003","+91 97458 97457",MOB,"","food@azadazad.com","azadspicegardens.com","","Company site + TradeIndia","https://azadazad.com/contact-us/","Heritage brand since 1940 - D2C store + recipe content can sell online, not just in stores"),
 ("Spice Brand","Kuruvi Spice","","Pappanamcode","Industrial Estate, Pappanamcode, TVM 695019","+91 80897 48927",MOB,"","contact@kuruvispice.com","kuruvispice.com","","Company website","https://www.kuruvispice.com/policies/contact-information","Young brand (Kuruvees Spices & Foods, inc. 2025) - growth stage, needs traffic + conversions"),
 ("Spice Brand","Spice Garden Foods (OPC) Pvt Ltd","G Pillai (CEO); Sajith (Owner contact)","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART + Facebook","https://www.indiamart.com/spicegardenfoods/","Only IndiaMART + Facebook page - a branded website builds trust with retailers"),
 ("Spice Brand","Badrinath Enterprises","","Karamana","TC 20/56 1, CITU Road, Melarannoor, Karamana PO, TVM 695002","","","","","","","IndiaMART","https://www.indiamart.com/badrinath-enterprises-thiruvananthapuram/spice-powder.html","Sambar/chilli/turmeric powders - catalogue site + WhatsApp ordering"),
 ("Spice Brand","North Malabar Spices Company","","Thumba","299/13, Near Kinfra, Thumba, TVM 695586","+91 98957 40107",MOB,"","","","","Spice Exporters Directory","https://spiceexportersdirectory.com/company/north-malabar-spices-company-thiruvananthapuram-wayanad-kerala-india/","Makes blended masalas - export catalogue site"),
 ("Spice Brand","Kadaas","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART","https://m.indiamart.com/kadaas/",""),
 ("Spice Brand","Henson Industries","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART","https://m.indiamart.com/henson-industries",""),
 ("Spice Brand","A J Spices","","Neyyattinkara","Neyyattinkara, Thiruvananthapuram","","","","","","","IndiaMART","https://m.indiamart.com/a-j-spices",""),
 ("Spice Brand","Bitepep","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","GlobalLinker","https://www.globallinker.com/seller/bitepep","Cardamom/bay leaf trader - B2B enquiry site"),
 ("Spice Brand","Wayanadan Spices","","Pattoor","Pattoor Junction, Thiruvananthapuram","","","","","","4.4 (214) Justdial","Justdial","https://www.justdial.com/Thiruvananthapuram/Masala-Powder-Manufacturers/nct-10313263","214 reviews but no website - online store would capture repeat buyers"),
 ("Spice Brand","Aishwarya Overseas Exporters","","Karamana","Karamana, Thiruvananthapuram","","","","","","","IndiaMART",S["im_spice"],"Exporter - export-buyer catalogue site"),
 ("Spice Brand","Devi Agencies","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART",S["im_spice"],""),
 # ---------------- FOOD ----------------
 ("Food Manufacturer","Nilamels & Kaimals Foods Pvt Ltd","","Menamkulam / Thumba","Plot A29(a), KINFRA Intl Apparel Park, Menamkulam, TVM","+91 471 310 2229",LAND,"+91 80438 32338 (IndiaMART)","","","","IndiaMART + Food Companies Directory","https://foodcompaniesdirectory.com/companies/nilamels-kaimals-foods-trivandrum-kerala-india/","Banana/tapioca chips manufacturer with no site - distributor & export catalogue"),
 ("Food Manufacturer","Mahasreevari Food Industries Pvt Ltd","","Kochuveli","Plot No 42, Industrial Development Area, Kochuveli, TVM","+91 80475 47662",IM,"","","","","IndiaMART",S["im_snack"],""),
 ("Food Manufacturer","Asian Food Products","","Pangappara","Pangappara, Thiruvananthapuram","+91 80477 65141",IM,"","","","","IndiaMART","https://www.indiamart.com/asian-food-products/","14 years on IndiaMART - own site cuts dependency on paid leads"),
 ("Food Manufacturer","Thankus Food Products","","Pappanamcode","Pappanamcode, Thiruvananthapuram","+91 80438 64029",IM,"","","","","IndiaMART",S["im_snack"],"GST-verified, 9 yrs on IndiaMART"),
 ("Food Manufacturer","T' Nacks Food Company","","Neyyattinkara","Maruthathoor, Tavaravila, Neyyattinkara","+91 80410 14452",IM,"","","","","IndiaMART",S["im_snack"],"Jackfruit chips - strong D2C/export story"),
 ("Food Manufacturer","SK Chips","","Parassala","Parassala, Thiruvananthapuram","+91 73565 51707",MOB,"","","skchips.in","","Company website","https://www.skchips.in/","Started 2024 - needs distributor enquiries fast"),
 ("Food Manufacturer","Dr Food Nutri Industries","","Kochuveli","Kochuveli, Thiruvananthapuram","+91 98953 04259",MOB,"+91 70344 85000","drfoodnutri@gmail.com","drfoodnutri.com","","Company website","https://www.drfoodnutri.com/","Site looks ~4 years old - refresh pitch"),
 ("Food Manufacturer","Tims Food Products","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART","https://dir.indiamart.com/thiruvananthapuram/snack-foods.html","13 yrs on IndiaMART"),
 ("Food Manufacturer","Jyothys Snacks","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART","https://www.indiamart.com/jyothys-snacks/snack-foods.html","Pickles + unniyappam - NRI/gifting online store angle"),
 ("Food Manufacturer","Bestree Foods","","Kadakkavoor","Kadakkavoor, Thiruvananthapuram","","","","","","","ExportersIndia","https://www.exportersindia.com/kerala/rice-flour.htm","Rice/idiyappam flour - supermarket & export catalogue"),
 ("Food Manufacturer","Trivandrum Flour Mills Pvt Ltd","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","Dun & Bradstreet",S["dnb_food"],""),
 ("Food Manufacturer","Crunchtopia (OPC) Pvt Ltd","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","Dun & Bradstreet",S["dnb_food"],""),
 ("Food Manufacturer","Cocoa and Crust Pvt Ltd","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","Dun & Bradstreet",S["dnb_food"],""),
 ("Food Manufacturer","Bread Factory / Azad Bakers & Confectioners","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART","https://www.indiamart.com/bread-factory/aboutus.html","Same group as Azad Spice - pitch both together"),
 ("Food Manufacturer","Divine Cakery","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART",S["im_bakery"],"Online cake ordering site"),
 ("Food Manufacturer","Veebakes","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART",S["im_bakery"],"Online ordering site"),
 ("Food Manufacturer","Mahendra Food Products","","Thiruvananthapuram","Thiruvananthapuram","","","","","","","IndiaMART",S["im_bakery"],""),
 ("Food Manufacturer","Jeevda Foods Pvt Ltd","","Jagathy","TC 44/580, Kochar Road, Jagathy, Thycaud PO, TVM 695014","","","","","","","InstaFinancials (MCA)","https://www.InstaFinancials.com/company/jeevda-foods-private-limited/U51900KL2019PTC056744",""),
 ("Food Manufacturer","Fresh N Fine Products (India) Pvt Ltd","","Menamkulam","46/47, KINFRA Apparel Park, Menamkulam, TVM","","","","","","","Zauba Corp (MCA)","https://www.zaubacorp.com/company-by-address/PLOT-NO-A39--KINFRA-PARK-MENAMKULAM--MENAMKULAM-P-O-----TRIVANDRUM--KERALA--INDIA---695582",""),
 # ---------------- INTERIOR ----------------
 ("Interior Design","Decon Infratech Pvt Ltd","","Trivandrum","Thiruvananthapuram","+91 75580 30104",MOB,"","","","","Archinect","https://archinect.com/deconinfratech","Architecture + interiors + build - turnkey project landing page"),
 ("Interior Design","Samson and Sons I Designs","","Muttada","TC 3/679, Kaliveena Bldg, TKD Rd, Muttada, TVM 695025","+91 471 393 9343",LAND,"","","","","in.worldorgs.com","https://in.worldorgs.com/catalog/thiruvananthapuram/interior-designer/idesigns",""),
 ("Interior Design","Havitive Architectural Studio (Havitive Infra Pvt Ltd)","","Kulathur / Kazhakoottam","Kulathur, Kazhakoottam, TVM","+91 92072 20320",MOB,"","","","","Web search listing","https://koloapp.in/kerala/flooring-experts/sajeev-s-8--thiruvananthapuram/posts/1628744734",""),
 ("Interior Design","Agnikone","","Kowdiar","Kowdiar, Thiruvananthapuram","+91 94474 24499",MOB,"","","","","GKW Retail directory",S["gkw"],"Homes + modular kitchens"),
 ("Interior Design","Atlas Interiors & Kitchen","","Kowdiar","Kowdiar, Thiruvananthapuram","+91 80780 90000",MOB,"","","atlasinteriors.in","","Website + directory","https://atlasinteriors.in/trivandrum/",""),
 ("Interior Design","Gazella Interiors","","Kowdiar","Kowdiar, Thiruvananthapuram","+91 75599 90009",MOB,"","info@gazellainteriors.com","gazellainteriors.com","","Company website","https://gazellainteriors.com/","Homes, offices, retail, hospitals - separate landing pages per segment"),
 ("Interior Design","Gruhem Interiors","","Pattom","Pattom, Thiruvananthapuram","+91 99953 96790",MOB,"","","","","GKW Retail directory",S["gkw"],""),
 ("Interior Design","Perfect Interiors","","Kowdiar","Opposite Kowdiar Palace, Kowdiar","","","","","","3.9 (53) Justdial","Justdial","https://www.justdial.com/Thiruvananthapuram/Perfect-Interiors-Opposite-Kowdiar-Palace-Kowdiar/0471PX471-X471-160601164311-E4E8_BZDET","Since 2010 - 15 years of projects, no portfolio site"),
 ("Interior Design","SDH Interiors & Exterior","","Pattom","Pattom, Thiruvananthapuram","08401486376",JD,"","","","","Justdial","https://www.justdial.com/Thiruvananthapuram/Interior-Designers-in-Pattom/nct-10272436",""),
 ("Interior Design","Des And Dec Pvt Ltd","","Kochulloor","Kochulloor, Thiruvananthapuram","08147128052",JD,"","","","","Justdial","https://www.justdial.com/Thiruvananthapuram/Interior-Designers-in-Kowdiar/nct-10272436",""),
 ("Interior Design","Kitchen Gallery","","Panavila","Panavila Junction, Thiruvananthapuram","08460495397",JD,"","","","","Justdial","https://www.justdial.com/Thiruvananthapuram/Interior-Designers-in-Pattom/nct-10272436","Modular kitchen - kitchen price estimator page"),
 ("Interior Design","De Lujan Decors","","Kazhakkoottam","Kazhakkoottam, Thiruvananthapuram","08460443680",JD,"","","","","Justdial",S["jd_int"],"Technopark IT crowd buying flats - target with a 2/3BHK package page"),
 ("Interior Design","Trendy Touches Interiors","","Kazhakkoottam","Kazhakkoottam, Thiruvananthapuram","","","","","","","Justdial",S["jd_int"],""),
 ("Interior Design","Vudpecker Modular Kitchens","","Sreekaryam","Sreekaryam showroom, Thiruvananthapuram","+91 95676 35659",MOB,"+91 95676 35521","info@vudpecker.com","vudpecker.com","","Company website","https://www.vudpecker.com/",""),
 ("Interior Design","Woodless Interiors","","Thycaud","Behind Central Residency, Thycaud, TVM","+91 97462 92910",MOB,"","projects@woodlessinteriors.com","woodlessinteriors.com","","Company website","https://woodlessinteriors.com/contacts/",""),
 ("Interior Design","Gud Interiors","","Nanthancode","MRA-19, Manchadivila Lane, Nanthancode, TVM","+91 70343 85555",MOB,"","info@gudinteriors.com","gudinteriors.com","","Company website","https://gudinteriors.com/",""),
 ("Interior Design","Spice Wood Interiors","","Kochuveli","Industrial Estate No 25, Near Titanium, Kochuveli, TVM 695021","+91 94464 43340",MOB,"","info@spicewood.in","spicewood.in","","Company website","https://www.spicewood.in/",""),
 ("Interior Design","Woodshine Interiors","","Trivandrum","Thiruvananthapuram","+91 80860 76666",MOB,"","woodshineinteriors@gmail.com","woodshine.in","","Company website","https://woodshine.in/contact-us/",""),
 ("Interior Design","DR Interio","","Enchakkal","Near Premiere Park Hotel, Enchakkal, TVM","+91 87141 13338",MOB,"","drinterio@gmail.com","drinterio.com","","Company website","https://drinterio.com/","Phone differs between page and metadata - confirm"),
 ("Interior Design","White Paper Interiors","","Poojappura","Poojappura, Thiruvananthapuram","+91 95678 62201",MOB,"","info@whitepaperinteriors.com","whitepaperinteriors.com","","Company website","https://whitepaperinteriors.com/",""),
 ("Interior Design","Fortune Interiors","","Sasthamangalam","Sasthamangalam, Thiruvananthapuram","+91 98470 62040",MOB,"","fortuneinteriorstvm@gmail.com","fortunexperience.com","","Company website","https://www.fortunexperience.com/services-interior-design-trivandrum/",""),
 ("Interior Design","VC Interiors","","NH-66","NH-66 Road, Thiruvananthapuram","+91 85906 22076",MOB,"","","vcinteriors.in","","Company website","https://vcinteriors.in/",""),
 ("Interior Design","BuildArt Interiors","","Trivandrum","Thiruvananthapuram","+91 81290 00120",MOB,"","info@buildartinterior.com","","","business-kerala.com directory","https://business-kerala.com/directory/listing/interior-designers-trivandrum","Website not confirmed - check before call"),
 # ---------------- FURNITURE ----------------
 ("Furniture Shop","Liya Furniture World","","Karamana","Karamana, Thiruvananthapuram","08511779243",JD,"","","","4.7 (27) Justdial","Justdial",S["jd_furn"],"New + second-hand furniture - 'sell your furniture' lead form"),
 ("Furniture Shop","Movable","","Anayara","NH Bypass Road, Anayara, TVM","09606280528",JD,"","","","4.1 (576) Justdial","Justdial",S["jd_furn"],"576 reviews - big footfall, no website to capture it online"),
 ("Furniture Shop","Dnm Furnishing","","Aryasala","Aryasala, Thiruvananthapuram","08401876235",JD,"","","","","Justdial (verified)",S["jd_furn"],""),
 ("Furniture Shop","New Star Agencies","","Karamana","Karamana, Thiruvananthapuram","08123232904",JD,"","","","","Justdial",S["jd_furn"],"Also does sofa repair - service booking form"),
 ("Furniture Shop","Rajashilpi Furniture & B S Timbers","","Kallara","Kallara, Thiruvananthapuram district","09845159549",JD,"","","","","Justdial",S["jd_furn"],"Wooden doors - catalogue for builders/contractors"),
 ("Furniture Shop","Naz Wood Industries","","Ambalathara","Thiruvallam, Ambalathara, TVM","08460281531",JD,"","","","","Justdial",S["jd_furn"],""),
 ("Furniture Shop","Antique Home Decor","","Poojappura","TC 17/1411(5), Gangothri Bldgs, near Chattambi Swami Temple, Poojappura","+91 98460 85350",MOB,"","","","","yappe.in","https://yappe.in/kerala/thiruvananthapuram/antique-home-decor/1710171","Antique furniture sells nationally - online catalogue reaches collectors"),
 ("Furniture Shop","Sofas & More by Stanley - Trivandrum","","Anayara","NH Bypass, Anayara, near Lords Hospital","+91 95443 56622",MOB,"","","","","yappe.in","https://yappe.in/kerala/thiruvananthapuram/sofas-more-by-stanley-trivandrum/1710175","Franchise - pitch local landing page + Google ads, not a full site"),
 ("Furniture Shop","Furniture Centre","","Karamana","Near Karamana Post Office, TVM","+91 471 234 9967",LAND,"","","","","yappe.in","https://yappe.in/kerala/thiruvananthapuram/furniture-centre/1710214",""),
 ("Furniture Shop","Thomson Furniture (Thomson Home Depot)","","Kesavadasapuram","Cordial Emerald, Kesavadasapuram-Ulloor Rd, TVM 695004","+91 95266 11221",MOB,"+91 95266 17733","thdcare@thomsonmails.com","thomsonfurniture.in","692 reviews (site)","Company website + Quickerala","https://thomsonfurniture.in/",""),
 ("Furniture Shop","Capito Furniture & Home Interiors","","Kaniyapuram","Kaniyapuram, near Kazhakkoottam","+91 99471 38000",MOB,"","","","","Quickerala",S["qk_furn"],"Large showroom - number from older listing, verify"),
 ("Furniture Shop","Popular Furniture","","Kazhakkoottam","Main Road (NH47), Kazhakkoottam","0471-2412761",LAND,"","","","","Quickerala","https://www.quickerala.com/trivandrum/kazhakkoottam/popular-furniture/84412","Listing ~9 years old - verify"),
 ("Furniture Shop","M M Furniture","","Kazhakkoottam","Bypass Road, opp Technopark main gate, Kazhakkoottam","","","","","","","Quickerala","https://www.quickerala.com/trivandrum/kazhakkoottam/m-m-furniture/427686","Opposite Technopark - target IT employees setting up homes"),
 ("Furniture Shop","Royal Agencies","","Balaramapuram","Balaramapuram, Thiruvananthapuram","0471-2400660",LAND,"","","","","Quickerala",S["qk_furn"],"One listing says closed - confirm before calling"),
 ("Furniture Shop","Diya Furniture","","Pongumoodu","Pongumoodu, Thiruvananthapuram","","","","","","","Justdial",S["jd_furn"],""),
 ("Furniture Shop","Sree Raja Furniture","","Nemom","Nemom, Thiruvananthapuram","","","","","","","Justdial",S["jd_furn"],"Also carpentry - custom furniture enquiry form"),
 ("Furniture Shop","A-One Furniture","","Kowdiar","Kowdiar, Thiruvananthapuram","","","","","","","Justdial",S["jd_furn"],""),
 ("Furniture Shop","Chellam Innova Furniture Palace","","Kowdiar","Kowdiar, Thiruvananthapuram","","","","","","","Justdial",S["jd_furn"],""),
 ("Furniture Shop","Subash Furniture","","Chalai","Chalai, Thiruvananthapuram","","","","","","","Justdial",S["jd_furn"],""),
 ("Furniture Shop","Nilambur Interiors & Furniture","","Nedumangad","Nedumangad, Thiruvananthapuram district","","","","","","","Quickerala",S["qk_furn"],"Manufacturer - factory-direct pricing site"),
 ("Furniture Shop","Ajantha Furniture","","Balaramapuram","Balaramapuram, Thiruvananthapuram","","","","","","","Quickerala",S["qk_furn"],""),
 ("Furniture Shop","Evernest Industries","","Nedumangad","Nedumangad, Thiruvananthapuram district","","","","","","","IndiaMART","https://dir.indiamart.com/thiruvananthapuram/wooden-furniture.html","Manufacturer/wholesaler"),
]

PITCH = {
 ("Solar Installation", False): "No website: launch a PM Surya Ghar / KSEB Soura subsidy landing page with a rooftop savings calculator and a WhatsApp 'Get free site survey' button, so homeowners searching 'solar subsidy Trivandrum' reach you directly instead of Justdial competitors.",
 ("Solar Installation", True): "Turn the site into a lead engine: subsidy + savings calculator, KSEB/MNRE empanelment proof, installed-project gallery (kW, savings/month), Google reviews and click-to-WhatsApp, with enquiry tracking so you can see cost per lead.",
 ("Spice Brand", False): "No website: branded D2C store with UPI checkout, WhatsApp ordering and a bulk/distributor enquiry page, so retailers and NRI buyers can find and order without going through IndiaMART.",
 ("Spice Brand", True): "Make the site sell: product pages with recipes, online checkout + Amazon/Flipkart links, FSSAI and quality proof, and separate retail / distributor / export enquiry forms.",
 ("Food Manufacturer", False): "No website: B2B catalogue site (product specs, MOQ, pack sizes, FSSAI, export readiness) with a distributor enquiry form, to win supermarket/export buyers and cut dependence on paid IndiaMART leads.",
 ("Food Manufacturer", True): "Upgrade to a distributor-lead funnel: downloadable catalogue, enquiry form with MOQ/location, retailer locator, export page and WhatsApp CTA; track enquiries per month.",
 ("Interior Design", False): "No website: portfolio site with project galleries, a 2BHK/3BHK cost estimator and a 'Book free consultation' form, plus Google reviews, so new homeowners shortlist you before calling competitors.",
 ("Interior Design", True): "Most interior sites show photos but don't capture leads: add a BHK cost estimator, before/after case studies with budgets, review widget, sticky WhatsApp CTA and a consultation booking calendar.",
 ("Furniture Shop", False): "No website: online catalogue with price ranges and an 'Enquire on WhatsApp' button per product, showroom map and offers page, so Instagram/Google browsers turn into showroom visits.",
 ("Furniture Shop", True): "Make the catalogue convert: price ranges, WhatsApp enquiry per product, festival-offer landing pages, EMI info and showroom-visit booking, linked to your Google Business profile.",
}

wb = Workbook()

# ======================= LEADS =======================
ws = wb.active
ws.title = "Leads"
ws.sheet_view.showGridLines = False
headers = ["Lead ID","Niche","Business Name","Owner / Decision Maker","Area","Address","Primary Phone",
           "Phone Type","WhatsApp","Alt / Office Phone","Email","Website","Website Status","Rating",
           "Lead Source","Source URL","Main Pitch","Priority","Lead Status","Calls Made","Last Call Date",
           "Next Follow-up","Deal Value (Rs)","Notes"]
widths = [9,17,32,26,16,34,17,22,22,22,30,26,24,16,22,28,60,10,22,10,13,13,14,40]
ws["A1"] = "Dezmos Digital Marketing - Trivandrum Website Leads"
ws["A1"].font = f_title
ws["A2"] = ("Niches: solar installation, spice brands, food manufacturers, interior design, furniture shops. "
            "Edit only the lime-tinted columns (Lead Status, Next Follow-up, Deal Value, Notes) and log every call in 'Call Log'. "
            "Calls Made and Last Call Date fill in automatically.")
ws["A2"].font = f_sub
ws.merge_cells("A2:X2")
ws.row_dimensions[2].height = 28
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
HR = 4
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=HR, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[HR].height = 30

def is_mobile(ph):
    d = "".join(ch for ch in ph if ch.isdigit())
    if d.startswith("91") and len(d) == 12: d = d[2:]
    return len(d) == 10 and d[0] in "6789"

first = HR + 1
for idx, r in enumerate(L):
    niche, name, owner, area, addr, ph, ptype, alt, email, web, rating, src, url, hook = r
    row = first + idx
    has_web = bool(web)
    if not has_web and email.endswith("@gmail.com"):
        wstatus = "No website (Gmail only)"
    elif not has_web:
        wstatus = NOWEB
    else:
        wstatus = WEB
    if ptype == MOB and is_mobile(ph):
        wa = ph + " (likely - verify)"
    elif ptype == JD:
        wa = "Use Justdial chat; ask for WhatsApp no."
    elif ptype == IM:
        wa = "Ask on first call"
    elif ptype == LAND:
        wa = alt if alt and is_mobile(alt.split(" (")[0]) else "Ask on first call"
    else:
        wa = ""
    pitch = PITCH[(niche, has_web)] + (f" Hook: {hook}." if hook else "")
    have_contact = bool(ph or email)
    if not has_web and have_contact and ptype in (MOB, LAND):
        prio = "Hot"
    elif have_contact:
        prio = "Warm"
    else:
        prio = "Cold"
    if not ptype and not ph:
        ptype = "Not public - find via website / visit"
    vals = [f"TVM-{idx+1:03d}", niche, name, owner or "Not listed - ask on call / MCA", area, addr, ph or "-",
            ptype, wa or "-", alt.strip() or "-", email or "-", web or "-", wstatus, rating or "-", src, url,
            pitch, prio, "Not Contacted",
            f"=COUNTIF('Call Log'!$B:$B,A{row})",
            f"=IF(T{row}=0,\"\",_xlfn.MAXIFS('Call Log'!$A:$A,'Call Log'!$B:$B,A{row}))",
            None, None, None]
    for ci, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=ci, value=v)
        c.font = f_body; c.border = border; c.alignment = wrap
        if idx % 2: c.fill = fill_gray
    ws.cell(row=row, column=16).hyperlink = url
    ws.cell(row=row, column=16).font = f_link
    if web:
        ws.cell(row=row, column=12).hyperlink = "https://" + web
        ws.cell(row=row, column=12).font = f_link
    for ci in (19, 22, 23, 24):
        ws.cell(row=row, column=ci).fill = fill_tint
    ws.cell(row=row, column=21).number_format = "dd-mmm-yy"
    ws.cell(row=row, column=22).number_format = "dd-mmm-yy"
    ws.cell(row=row, column=23).number_format = "#,##0;(#,##0);-"
    ws.cell(row=row, column=20).alignment = center
    ws.cell(row=row, column=18).alignment = center
last = first + len(L) - 1
MAXROW = 400  # room to add more leads

dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
dv.add(f"S{first}:S{MAXROW}"); ws.add_data_validation(dv)
dvp = DataValidation(type="list", formula1='"Hot,Warm,Cold"', allow_blank=True)
dvp.add(f"R{first}:R{MAXROW}"); ws.add_data_validation(dvp)
dvn = DataValidation(type="list", formula1='"' + ",".join(NICHES) + '"', allow_blank=True)
dvn.add(f"B{first}:B{MAXROW}"); ws.add_data_validation(dvn)

ws.conditional_formatting.add(f"R{first}:R{MAXROW}", CellIsRule(operator="equal", formula=['"Hot"'], fill=fill_lime, font=Font(name=FONT, bold=True, color=FOREST)))
ws.conditional_formatting.add(f"S{first}:S{MAXROW}", CellIsRule(operator="equal", formula=['"Converted"'], fill=fill_lime, font=Font(name=FONT, bold=True, color=FOREST)))
ws.conditional_formatting.add(f"M{first}:M{MAXROW}", FormulaRule(formula=[f'LEFT(M{first},2)="No"'], font=Font(name=FONT, bold=True, color=FOREST), fill=fill_tint))
ws.freeze_panes = "D5"
ws.auto_filter.ref = f"A{HR}:X{last}"
ws["H4"].comment = Comment("Justdial / IndiaMART virtual numbers forward to the business. Ask for the owner's direct mobile and WhatsApp on the first call and update the row.", "Dezmos")
ws["I4"].comment = Comment("WhatsApp shown as 'likely' when the listed number is a 10-digit mobile. Confirm with a WhatsApp check before messaging.", "Dezmos")
ws["M4"].comment = Comment("'No website found' = none found in directory/web search results. 'Has website' sites could not be opened from the research environment, so audit them for lead capture (form, WhatsApp CTA, speed, mobile, reviews) before pitching.", "Dezmos")

# ======================= CALL LOG =======================
cl = wb.create_sheet("Call Log")
cl.sheet_view.showGridLines = False
cl["A1"] = "Call Log - one row per call"
cl["A1"].font = f_title
cl["A2"] = "Enter Date, Lead ID (e.g. TVM-001), Outcome and notes. Business and Niche fill in automatically. The grey row 4 is an example and is not counted (its ID does not match any lead)."
cl["A2"].font = f_sub
cl.merge_cells("A2:I2")
ch = ["Date","Lead ID","Business (auto)","Niche (auto)","Call Outcome","Spoke With","Next Step","Follow-up Date","Notes"]
cw = [13,11,32,18,22,20,30,14,40]
for i, (h, w) in enumerate(zip(ch, cw), 1):
    c = cl.cell(row=3, column=i, value=h)
    c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
    cl.column_dimensions[c.column_letter].width = w
ex = ["2026-10-21", "EXAMPLE", "Example Solar Pvt Ltd", "Solar Installation", "Call Back Later", "Owner (Mr. Name)", "Send WhatsApp portfolio", "2026-10-23", "Asked to call after 5 pm"]
from datetime import date
ex[0] = date(2026, 10, 21); ex[7] = date(2026, 10, 23)
for i, v in enumerate(ex, 1):
    c = cl.cell(row=4, column=i, value=v)
    c.font = Font(name=FONT, size=10, italic=True, color="8A9294"); c.fill = fill_gray; c.border = border
for r in range(5, 1001):
    cl.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(INDEX(Leads!$C:$C,MATCH(B{r},Leads!$A:$A,0)),"ID not found"))')
    cl.cell(row=r, column=4, value=f'=IF(B{r}="","",IFERROR(INDEX(Leads!$B:$B,MATCH(B{r},Leads!$A:$A,0)),""))')
    for ci in range(1, 10):
        c = cl.cell(row=r, column=ci); c.font = f_body; c.border = border
    cl.cell(row=r, column=1).number_format = "dd-mmm-yy"
    cl.cell(row=r, column=8).number_format = "dd-mmm-yy"
    for ci in (1, 2, 5, 6, 7, 8, 9):
        cl.cell(row=r, column=ci).fill = fill_tint
cl.cell(row=4, column=1).number_format = "dd-mmm-yy"; cl.cell(row=4, column=8).number_format = "dd-mmm-yy"
dvo = DataValidation(type="list", formula1='"' + ",".join(OUTCOMES) + '"', allow_blank=True)
dvo.add("E4:E1000"); cl.add_data_validation(dvo)
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
dvd.add("A5:A1000"); dvd.add("H5:H1000"); cl.add_data_validation(dvd)
cl.freeze_panes = "A4"

# ======================= DASHBOARD =======================
db = wb.create_sheet("Dashboard", 0)
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMN", [24,14,14,14,14,14,3,24,14,3,14,14,14,14]):
    db.column_dimensions[col].width = w
db["A1"] = "Lead Pipeline Dashboard"
db["A1"].font = f_title
db["A2"] = "Updates automatically from the Leads and Call Log sheets. Charts fill in as you log calls and set Lead Status."
db["A2"].font = f_sub

LR = f"Leads!$A${first}:$A${MAXROW}"
LN = f"Leads!$B${first}:$B${MAXROW}"
LS = f"Leads!$S${first}:$S${MAXROW}"
LC = f"Leads!$T${first}:$T${MAXROW}"
LM = f"Leads!$M${first}:$M${MAXROW}"
LV = f"Leads!$W${first}:$W${MAXROW}"
CO = "'Call Log'!$E$5:$E$1000"
CN = "'Call Log'!$D$5:$D$1000"

kpis = [
    ("Total Leads", f"=COUNTA({LR})", "0"),
    ("Calls Made", f"=SUM({LC})", "0"),
    ("Leads Contacted", f'=COUNTIF({LC},">0")', "0"),
    ("Interested / Meeting", f'=COUNTIF({LS},"Interested")+COUNTIF({LS},"Meeting / Proposal Sent")', "0"),
    ("Converted", f'=COUNTIF({LS},"Converted")', "0"),
    ("Conversion Rate", "=IFERROR(E5/C5,0)", "0.0%"),
]
# KPI cards across A..F row 4/5
for i, (lab, f, fmt) in enumerate(kpis):
    col = i + 1
    t = db.cell(row=4, column=col, value=lab)
    t.font = Font(name=FONT, size=9, bold=True, color=MUTED); t.alignment = center; t.fill = fill_gray
    v = db.cell(row=5, column=col, value=f)
    v.font = f_kpi; v.alignment = center; v.number_format = fmt
    v.fill = fill_lime if lab == "Converted" else fill_gray
db.row_dimensions[5].height = 40
db["A6"] = "Conversion rate = Converted / Leads Contacted."
db["A6"].font = f_sub
db["H4"] = "Revenue Won (Rs)"; db["H4"].font = Font(name=FONT, size=9, bold=True, color=MUTED); db["H4"].fill = fill_gray; db["H4"].alignment = center
db["H5"] = f'=SUMIFS({LV},{LS},"Converted")'; db["H5"].font = f_kpi; db["H5"].fill = fill_gray; db["H5"].alignment = center; db["H5"].number_format = "#,##0;(#,##0);-"
db["I4"] = "No-Website Leads"; db["I4"].font = Font(name=FONT, size=9, bold=True, color=MUTED); db["I4"].fill = fill_gray; db["I4"].alignment = center
db["I5"] = f'=COUNTIF({LM},"No website*")'; db["I5"].font = f_kpi; db["I5"].fill = fill_gray; db["I5"].alignment = center

def head(row, cols, start=1):
    for i, h in enumerate(cols):
        c = db.cell(row=row, column=start + i, value=h)
        c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border

# By niche table
db["A8"] = "Performance by Niche"; db["A8"].font = f_bold
head(9, ["Niche", "Leads", "Calls Made", "Contacted", "Converted", "Conv. Rate"])
for i, n in enumerate(NICHES):
    r = 10 + i
    db.cell(row=r, column=1, value=n)
    db.cell(row=r, column=2, value=f'=COUNTIF({LN},A{r})')
    db.cell(row=r, column=3, value=f'=SUMIFS({LC},{LN},A{r})')
    db.cell(row=r, column=4, value=f'=COUNTIFS({LN},A{r},{LC},">0")')
    db.cell(row=r, column=5, value=f'=COUNTIFS({LN},A{r},{LS},"Converted")')
    db.cell(row=r, column=6, value=f'=IFERROR(E{r}/D{r},0)')
    for c in range(1, 7):
        cell = db.cell(row=r, column=c); cell.font = f_body; cell.border = border
        if i % 2: cell.fill = fill_gray
    db.cell(row=r, column=6).number_format = "0.0%"
r = 15
db.cell(row=r, column=1, value="Total")
for c, L_ in zip(range(2, 6), "BCDE"):
    db.cell(row=r, column=c, value=f"=SUM({L_}10:{L_}14)")
db.cell(row=r, column=6, value="=IFERROR(E15/D15,0)").number_format = "0.0%"
for c in range(1, 7):
    db.cell(row=r, column=c).font = f_bold; db.cell(row=r, column=c).border = border; db.cell(row=r, column=c).fill = fill_tint

# Call outcomes table (per call)
db["H8"] = "Call Outcomes (all calls)"; db["H8"].font = f_bold
head(9, ["Outcome", "Calls"], start=8)
for i, o in enumerate(OUTCOMES):
    rr = 10 + i
    db.cell(row=rr, column=8, value=o).font = f_body
    db.cell(row=rr, column=9, value=f'=COUNTIF({CO},H{rr})').font = f_body
    for c in (8, 9): db.cell(row=rr, column=c).border = border
db.cell(row=18, column=8, value="Total").font = f_bold
db.cell(row=18, column=9, value="=SUM(I10:I17)").font = f_bold

# Conversion split table (contacted leads)
db["K8"] = "Contacted Leads: Converted vs Not"; db["K8"].font = f_bold
head(9, ["Result", "Leads"], start=11)
db["K10"] = "Converted"; db["L10"] = "=E5"
db["K11"] = "Not converted yet"; db["L11"] = "=MAX(C5-E5,0)"
db["K12"] = "Not contacted"; db["L12"] = "=MAX(A5-C5,0)"
for rr in (10, 11, 12):
    for c in (11, 12):
        db.cell(row=rr, column=c).font = f_body; db.cell(row=rr, column=c).border = border

# Lead status table
db["K14"] = "Lead Status"; db["K14"].font = f_bold
head(15, ["Status", "Leads"], start=11)
for i, s in enumerate(STATUSES):
    rr = 16 + i
    db.cell(row=rr, column=11, value=s).font = f_body
    db.cell(row=rr, column=12, value=f'=COUNTIF({LS},K{rr})').font = f_body
    for c in (11, 12): db.cell(row=rr, column=c).border = border

PALETTE = [FOREST, LIME, "6E7B7F", "A9D96B", "C9CFD1", "3F5258", "E2FFC0", "8FA0A4"]

def pie(title, labels_ref, data_ref, anchor, colors=PALETTE):
    p = PieChart()
    p.title = title
    p.add_data(data_ref, titles_from_data=False)
    p.set_categories(labels_ref)
    p.dataLabels = DataLabelList(); p.dataLabels.showPercent = True
    p.dataLabels.showVal = False; p.dataLabels.showCatName = False; p.dataLabels.showSerName = False; p.dataLabels.showLegendKey = False; p.dataLabels.showLeaderLines = False
    p.height, p.width = 8.5, 12
    s = p.series[0]
    n = labels_ref.max_row - labels_ref.min_row + 1
    for i in range(n):
        pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = colors[i % len(colors)]
        pt.graphicalProperties.line.solidFill = "FFFFFF"
        s.dPt.append(pt)
    db.add_chart(p, anchor)

pie("Call Outcomes", Reference(db, min_col=8, min_row=10, max_row=17), Reference(db, min_col=9, min_row=10, max_row=17), "A26")
pie("Converted vs Not Converted", Reference(db, min_col=11, min_row=10, max_row=12), Reference(db, min_col=12, min_row=10, max_row=12), "F26", [LIME, FOREST, "C9CFD1"])
pie("Leads by Niche", Reference(db, min_col=1, min_row=10, max_row=14), Reference(db, min_col=2, min_row=10, max_row=14), "A44")

bar = BarChart(); bar.type = "col"; bar.title = "Calls vs Conversions by Niche"
bar.add_data(Reference(db, min_col=3, min_row=9, max_row=14), titles_from_data=True)
bar.add_data(Reference(db, min_col=5, min_row=9, max_row=14), titles_from_data=True)
bar.set_categories(Reference(db, min_col=1, min_row=10, max_row=14))
bar.series[0].graphicalProperties.solidFill = FOREST
bar.series[1].graphicalProperties.solidFill = LIME
bar.height, bar.width = 8.5, 12
db.add_chart(bar, "F44")
db["A25"] = "Charts"; db["A25"].font = f_bold

# ======================= PITCH PLAYBOOK =======================
pb = wb.create_sheet("Pitch Playbook")
pb.sheet_view.showGridLines = False
pb["A1"] = "Pitch Playbook - website development"
pb["A1"].font = f_title
pb.column_dimensions["A"].width = 20
for col in "BCDE": pb.column_dimensions[col].width = 48
head_cols = ["Niche", "Core pitch (no website)", "Core pitch (has website, not converting)", "WhatsApp / call opener", "Objection -> answer"]
for i, h in enumerate(head_cols, 1):
    c = pb.cell(row=3, column=i, value=h); c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
OPENERS = {
 "Solar Installation": ("Hi [Name], this is [Your Name] from Dezmos Digital Marketing, Trivandrum. People here search 'solar subsidy Trivandrum' every day. We build solar websites with a subsidy calculator that send enquiries straight to your WhatsApp. Can I show you a 2-minute demo?",
                        "'We get leads from Justdial' -> Justdial shares each lead with 3-5 competitors. Your own site gives you leads nobody else gets."),
 "Spice Brand": ("Hi [Name], [Your Name] from Dezmos, Trivandrum. Your masalas could sell online to NRI and pan-India buyers. We set up spice brand stores with UPI checkout and WhatsApp ordering. Open to a quick look?",
                 "'Distributors handle sales' -> A site supports distributors with a catalogue, and adds direct online orders with better margins."),
 "Food Manufacturer": ("Hi [Name], [Your Name] from Dezmos. Supermarket and export buyers check a manufacturer's website before placing orders. We build B2B catalogue sites with distributor enquiry forms. Could I send a sample?",
                       "'IndiaMART is enough' -> IndiaMART leads are paid and shared. A site plus Google listing brings direct buyers at no per-lead cost."),
 "Interior Design": ("Hi [Name], [Your Name] from Dezmos, Trivandrum. Homeowners shortlist interior firms online before calling. We build portfolio sites with a BHK cost estimator that books consultations for you. Can I show you one?",
                     "'Instagram is enough' -> Instagram shows work; a site with an estimator and booking form captures the enquiry and ranks on Google."),
 "Furniture Shop": ("Hi [Name], [Your Name] from Dezmos. Buyers browse furniture online before visiting a showroom. We build catalogue sites with a WhatsApp enquiry button on every product. Can I share an example?",
                    "'Customers walk in anyway' -> Most buyers check online first. A catalogue sends them to your showroom instead of the chain stores."),
}
for i, n in enumerate(NICHES):
    r = 4 + i
    vals = [n, PITCH[(n, False)], PITCH[(n, True)], OPENERS[n][0], OPENERS[n][1]]
    for ci, v in enumerate(vals, 1):
        c = pb.cell(row=r, column=ci, value=v); c.font = f_body if ci > 1 else f_bold; c.alignment = wrap; c.border = border
        if i % 2: c.fill = fill_gray
    pb.row_dimensions[r].height = 120
pb["A10"] = "Website audit checklist (before pitching a 'has website' lead)"; pb["A10"].font = f_bold
checks = ["Opens fast on mobile (under 3 seconds)?", "Click-to-call and WhatsApp button visible on every page?",
          "Enquiry form above the fold, short (3-4 fields)?", "Proof shown: reviews, certifications, project photos with numbers?",
          "Google Business profile linked and consistent (name, phone, address)?", "Ranks for '[service] Trivandrum' on Google?",
          "SSL (https) and up-to-date content (no old offers / dead links)?", "Any enquiry tracking (Analytics, form notifications)?"]
for i, t in enumerate(checks):
    pb.cell(row=11 + i, column=1, value=f"{i+1}.").font = f_body
    c = pb.cell(row=11 + i, column=2, value=t); c.font = f_body
pb["A20"] = "Compliance note: these are publicly listed business contacts. Call in business hours, honour 'do not contact' requests, and avoid bulk unsolicited WhatsApp messages (TRAI / WhatsApp Business policy)."
pb["A20"].font = f_sub
pb.merge_cells("A20:E20")

# ======================= SOURCES =======================
so = wb.create_sheet("Sources & Notes")
so.sheet_view.showGridLines = False
so.column_dimensions["A"].width = 110
lines = [
 ("How these leads were collected", True),
 ("Researched on 20 Oct 2026 from public business listings: Justdial, IndiaMART, TradeIndia, ExportersIndia, ENF Solar, KSEB Soura empanelled list, Quickerala, yappe.in, Dun & Bradstreet, MCA-based registries and company websites.", False),
 ("Google Maps (GMB) scraping was not possible from the research environment (Google was blocked by the network policy), so GMB-only businesses are not in this list yet.", False),
 ("Owner names are rarely public. Where none was found, the column says 'Not listed - ask on call / MCA'. Check directors on the MCA portal (mca.gov.in) for Pvt Ltd companies.", False),
 ("Justdial (08xxx / 09xxx 11-digit) and IndiaMART (+91 80xxx) numbers are virtual numbers that forward to the business. Ask for the direct mobile and WhatsApp on the first call.", False),
 ("Directory data goes stale. Verify each number before relying on it; some listings are several years old (noted in the Pitch hook / Notes).", False),
 ("National chains (Pepperfry, Wakefit, Royaloak, Damro, Livspace, D'LIFE) were left out on purpose: they have corporate websites and central marketing teams.", False),
 ("", False),
 ("Priority logic", True),
 ("Hot = no website found + direct mobile or landline.  Warm = has contact details (website already, or reachable only via virtual number).  Cold = no public phone/email yet.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = so.cell(row=i, column=1, value=t); c.font = f_bold if b else f_body; c.alignment = Alignment(wrap_text=True)

for sh in wb.worksheets:
    sh.sheet_properties.tabColor = LIME if sh.title == "Dashboard" else FOREST

wb.save(OUT)
print("leads:", len(L))
